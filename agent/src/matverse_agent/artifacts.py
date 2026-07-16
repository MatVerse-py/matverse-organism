from __future__ import annotations

import html
import json
import re
import shutil
import subprocess
import textwrap
from pathlib import Path
from typing import Any


def _slug(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    if not slug:
        raise ValueError("Name must contain at least one alphanumeric character")
    return slug[:80]


def _write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


def create_website(
    workspace: Path,
    name: str,
    title: str,
    sections: list[dict[str, str]],
) -> dict[str, Any]:
    root = workspace / "artifacts" / "websites" / _slug(name)
    section_html = "\n".join(
        f"<section><h2>{html.escape(item.get('heading', ''))}</h2>"
        f"<p>{html.escape(item.get('body', ''))}</p></section>"
        for item in sections
    )
    index = f"""<!doctype html>
<html lang="pt-BR">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <meta name="description" content="{html.escape(title)}">
  <title>{html.escape(title)}</title>
  <link rel="stylesheet" href="styles.css">
</head>
<body>
  <header><span class="mark">MATVERSE</span><h1>{html.escape(title)}</h1></header>
  <main>{section_html}</main>
  <footer>Gerado localmente pelo MatVerse Agent</footer>
  <script src="app.js"></script>
</body>
</html>
"""
    css = """:root{color-scheme:dark;--bg:#090b10;--panel:#131722;--text:#f5f7ff;--muted:#aeb7ca;--accent:#77f2c1}*{box-sizing:border-box}body{margin:0;background:radial-gradient(circle at top,#182033,var(--bg) 45%);color:var(--text);font:18px/1.6 system-ui,sans-serif}header,main,footer{width:min(100% - 2rem,960px);margin:auto}header{padding:7rem 0 3rem}.mark{color:var(--accent);font-weight:800;letter-spacing:.2em}h1{font-size:clamp(2.5rem,8vw,6rem);line-height:.95;margin:.5rem 0}main{display:grid;gap:1rem;padding-bottom:5rem}section{background:color-mix(in srgb,var(--panel) 90%,transparent);border:1px solid #293149;border-radius:20px;padding:2rem;box-shadow:0 20px 60px #0005}h2{margin-top:0;color:var(--accent)}p{color:var(--muted)}footer{border-top:1px solid #293149;padding:2rem 0;color:var(--muted)}@media(max-width:600px){header{padding-top:4rem}section{padding:1.25rem}}"""
    js = "document.documentElement.dataset.ready='true';\n"
    _write(root / "index.html", index)
    _write(root / "styles.css", css)
    _write(root / "app.js", js)
    return {"status": "PASS", "kind": "website", "path": str(root)}


def create_mobile_pwa(
    workspace: Path,
    name: str,
    title: str,
    features: list[str],
) -> dict[str, Any]:
    root = workspace / "artifacts" / "mobile" / _slug(name)
    cards = "\n".join(
        f"<li><strong>{index + 1:02d}</strong><span>{html.escape(feature)}</span></li>"
        for index, feature in enumerate(features)
    )
    index = f"""<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover"><meta name="theme-color" content="#0b1020"><link rel="manifest" href="manifest.webmanifest"><link rel="stylesheet" href="styles.css"><title>{html.escape(title)}</title></head><body><main><p class="eyebrow">MATVERSE LOCAL APP</p><h1>{html.escape(title)}</h1><ul>{cards}</ul><button id="install" hidden>Instalar aplicativo</button></main><script src="app.js"></script></body></html>"""
    css = """*{box-sizing:border-box}body{margin:0;min-height:100dvh;background:#0b1020;color:#f8faff;font:17px/1.5 system-ui,sans-serif;padding:env(safe-area-inset-top) 1rem env(safe-area-inset-bottom)}main{max-width:680px;margin:auto;padding:5rem 0}.eyebrow{color:#5ff0bd;font-weight:800;letter-spacing:.18em}h1{font-size:clamp(3rem,13vw,6rem);line-height:.95}ul{list-style:none;padding:0;display:grid;gap:.8rem}li{display:flex;gap:1rem;padding:1.2rem;background:#151d35;border:1px solid #2b385e;border-radius:18px}li strong{color:#5ff0bd}button{width:100%;padding:1rem;border:0;border-radius:14px;background:#5ff0bd;color:#07120f;font-weight:800;font-size:1rem}"""
    app_js = """if('serviceWorker'in navigator){navigator.serviceWorker.register('./sw.js')}let promptEvent;const button=document.querySelector('#install');window.addEventListener('beforeinstallprompt',event=>{event.preventDefault();promptEvent=event;button.hidden=false});button.addEventListener('click',async()=>{if(promptEvent){promptEvent.prompt();await promptEvent.userChoice;button.hidden=true}});"""
    manifest = {
        "name": title,
        "short_name": title[:24],
        "start_url": "./",
        "display": "standalone",
        "background_color": "#0b1020",
        "theme_color": "#0b1020",
        "icons": [],
    }
    sw = """const CACHE='matverse-v1';const ASSETS=['./','./index.html','./styles.css','./app.js','./manifest.webmanifest'];self.addEventListener('install',event=>event.waitUntil(caches.open(CACHE).then(cache=>cache.addAll(ASSETS))));self.addEventListener('fetch',event=>event.respondWith(caches.match(event.request).then(hit=>hit||fetch(event.request))));"""
    _write(root / "index.html", index)
    _write(root / "styles.css", css)
    _write(root / "app.js", app_js)
    _write(root / "manifest.webmanifest", json.dumps(manifest, ensure_ascii=False, indent=2))
    _write(root / "sw.js", sw)
    return {"status": "PASS", "kind": "installable_pwa", "path": str(root)}


def create_slides(
    workspace: Path,
    name: str,
    title: str,
    slides: list[dict[str, Any]],
) -> dict[str, Any]:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:
        return {"status": "HOLD", "reason": f"Install artifact extras: {exc}"}

    root = workspace / "artifacts" / "slides"
    root.mkdir(parents=True, exist_ok=True)
    output = root / f"{_slug(name)}.pptx"
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    first = presentation.slides.add_slide(presentation.slide_layouts[0])
    first.shapes.title.text = title
    first.placeholders[1].text = "Gerado localmente pelo MatVerse Agent"

    for item in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = str(item.get("title", ""))
        body = slide.placeholders[1].text_frame
        body.clear()
        bullets = item.get("bullets", [])
        if not isinstance(bullets, list):
            bullets = [str(bullets)]
        for index, bullet in enumerate(bullets):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = str(bullet)
            paragraph.font.size = Pt(24)
    presentation.save(output)
    return {"status": "PASS", "kind": "pptx", "path": str(output), "slides": len(slides) + 1}


def create_video(
    workspace: Path,
    name: str,
    scenes: list[dict[str, Any]],
    width: int = 1280,
    height: int = 720,
    fps: int = 30,
) -> dict[str, Any]:
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError as exc:
        return {"status": "HOLD", "reason": f"Install artifact extras: {exc}"}
    if shutil.which("ffmpeg") is None:
        return {"status": "HOLD", "reason": "ffmpeg executable not found"}
    if not scenes:
        raise ValueError("At least one scene is required")

    root = workspace / "artifacts" / "video" / _slug(name)
    frames = root / "frames"
    frames.mkdir(parents=True, exist_ok=True)
    font = ImageFont.load_default()
    concat_lines: list[str] = []

    for index, scene in enumerate(scenes):
        image = Image.new("RGB", (width, height), "#090b10")
        draw = ImageDraw.Draw(image)
        heading = str(scene.get("title", f"Cena {index + 1}"))
        body = str(scene.get("body", ""))
        duration = float(scene.get("duration", 3.0))
        draw.rounded_rectangle((70, 70, width - 70, height - 70), radius=35, fill="#151d35", outline="#5ff0bd", width=4)
        draw.text((120, 130), heading, fill="#5ff0bd", font=font)
        wrapped = "\n".join(textwrap.wrap(body, width=72))
        draw.multiline_text((120, 220), wrapped, fill="#f8faff", font=font, spacing=12)
        frame = frames / f"{index:04d}.png"
        image.save(frame)
        concat_lines.extend([f"file '{frame.as_posix()}'", f"duration {duration:.3f}"])
    concat_lines.append(f"file '{(frames / f'{len(scenes)-1:04d}.png').as_posix()}'")
    concat_file = root / "concat.txt"
    _write(concat_file, "\n".join(concat_lines) + "\n")
    output = root / f"{_slug(name)}.mp4"
    process = subprocess.run(
        [
            "ffmpeg",
            "-y",
            "-f",
            "concat",
            "-safe",
            "0",
            "-i",
            str(concat_file),
            "-vf",
            f"fps={fps},format=yuv420p",
            "-movflags",
            "+faststart",
            str(output),
        ],
        capture_output=True,
        text=True,
        timeout=300,
        check=False,
    )
    if process.returncode != 0:
        return {"status": "BLOCK", "reason": process.stderr[-4000:]}
    return {"status": "PASS", "kind": "mp4", "path": str(output), "scenes": len(scenes)}
